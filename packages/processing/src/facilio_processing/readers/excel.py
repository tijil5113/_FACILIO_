"""Excel .xlsx reader.

Formulas are not executed. Cached values (when present) are read as data.
Multiple usable sheets require an explicit selection.
"""

from __future__ import annotations

from dataclasses import asdict
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from facilio_processing.errors import (
    EmptyDatasetError,
    InvalidSheetError,
    InvalidWorkbookError,
    SheetSelectionRequiredError,
)
from facilio_processing.inference import columns_from_frame
from facilio_processing.serialization import frame_preview_rows
from facilio_processing.types import PreviewResult, ReadResult, SheetInfo, WorkbookInspection


class ExcelReader:
    def inspect(self, path: Path) -> WorkbookInspection:
        workbook = _open_workbook(path)
        try:
            sheets = [_describe_sheet(workbook[name]) for name in workbook.sheetnames]
            return WorkbookInspection(sheets=sheets)
        finally:
            workbook.close()

    def read(self, path: Path, *, sheet: str | None = None) -> ReadResult:
        selected = self._resolve_sheet(path, sheet)
        frame = _load_sheet(path, selected)
        columns = columns_from_frame(frame)
        return ReadResult(
            frame=frame,
            columns=columns,
            row_count=int(frame.shape[0]),
            column_count=int(frame.shape[1]),
            selected_sheet=selected,
        )

    def preview(
        self,
        path: Path,
        *,
        max_rows: int,
        max_columns: int,
        sheet: str | None = None,
    ) -> PreviewResult:
        result = self.read(path, sheet=sheet)
        rows, truncated_rows, truncated_columns = frame_preview_rows(
            result.frame, max_rows=max_rows, max_columns=max_columns
        )
        return PreviewResult(
            columns=result.columns[:max_columns],
            rows=rows,
            row_count=result.row_count,
            column_count=result.column_count,
            preview_row_count=len(rows),
            truncated_rows=truncated_rows,
            truncated_columns=truncated_columns,
            selected_sheet=result.selected_sheet,
        )

    def _resolve_sheet(self, path: Path, sheet: str | None) -> str:
        inspection = self.inspect(path)
        usable = inspection.usable
        if sheet is not None:
            match = next((item for item in inspection.sheets if item.name == sheet), None)
            if match is None:
                raise InvalidSheetError(f'Sheet "{sheet}" was not found in the workbook.')
            if match.empty:
                raise InvalidSheetError(f'Sheet "{sheet}" does not contain tabular data.')
            return sheet
        if not usable:
            raise EmptyDatasetError("The workbook does not contain a usable sheet.")
        if len(usable) == 1:
            return usable[0].name
        raise SheetSelectionRequiredError(
            "This workbook has multiple sheets. Select one to ingest.",
            details={"sheets": [asdict(item) for item in inspection.sheets]},
        )


def _open_workbook(path: Path):
    try:
        return load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError) as error:
        raise InvalidWorkbookError("The file is not a valid Excel workbook.") from error


def _describe_sheet(worksheet: Worksheet) -> SheetInfo:
    hidden = getattr(worksheet, "sheet_state", "visible") != "visible"
    return SheetInfo(name=worksheet.title, empty=_sheet_is_empty(worksheet), hidden=hidden)


def _sheet_is_empty(worksheet: Worksheet) -> bool:
    for row in worksheet.iter_rows(values_only=True):
        if any(cell is not None and str(cell).strip() != "" for cell in row):
            return False
    return True


def _load_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    workbook = _open_workbook(path)
    try:
        if sheet_name not in workbook.sheetnames:
            raise InvalidSheetError(
                f'Sheet "{sheet_name}" was not found in the workbook.'
            )
        worksheet = workbook[sheet_name]
        rows = [
            [_cell(cell) for cell in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()

    while rows and all(cell is None for cell in rows[-1]):
        rows.pop()
    if not rows:
        raise EmptyDatasetError(f'Sheet "{sheet_name}" does not contain tabular data.')

    width = max(len(row) for row in rows)
    normalized = [row + [None] * (width - len(row)) for row in rows]
    headers = [
        "" if cell is None else str(cell) for cell in normalized[0]
    ]
    body = normalized[1:]
    if not any(name.strip() for name in headers) and not body:
        raise EmptyDatasetError(f'Sheet "{sheet_name}" does not contain tabular data.')
    frame = pd.DataFrame(body, columns=headers, dtype=object)
    return frame


def _cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str) and value == "":
        return None
    return value
